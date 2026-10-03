import os
import shutil
from datetime import datetime
from openpyxl import Workbook, load_workbook
from . import config

COLUMNS = ["timestamp", "customer_name", "customer_id", "product_names", "qty", "total_price", "address", "status"]


def _path(bot_id: int) -> str:
    d = os.path.join(config.UPLOAD_DIR, str(bot_id))
    os.makedirs(d, exist_ok=True)
    return os.path.join(d, "orders.xlsx")


def _rotate_backup(bot_id: int):
    path = _path(bot_id)
    if not os.path.exists(path):
        return
    stamp = datetime.now().strftime("%Y%m%d%H%M%S")
    shutil.copy2(path, f"{path}.{stamp}.bak")
    backups = sorted(
        f for f in os.listdir(os.path.dirname(path))
        if f.startswith("orders.xlsx.") and f.endswith(".bak")
    )
    for old in backups[:-5]:
        os.remove(os.path.join(os.path.dirname(path), old))


def append_order(bot_id: int, order: dict, customer_user_id: str, status: str = "pending"):
    path = _path(bot_id)
    if os.path.exists(path):
        _rotate_backup(bot_id)
        wb = load_workbook(path)
        ws = wb.active
    else:
        wb = Workbook()
        ws = wb.active
        ws.title = "orders"
        ws.append(COLUMNS)
    products = order.get("products") or []
    names = ", ".join(str(p.get("product_name", "")) for p in products)
    qty = sum(int(p.get("quantity", 0) or 0) for p in products)
    ws.append([
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        order.get("customer_name") or "",
        customer_user_id,
        names,
        qty,
        order.get("total_price"),
        order.get("delivery_address") or "",
        status,
    ])
    wb.save(path)
