import asyncio
import hmac
import json
import os
import queue
import re
import threading
import time
from types import SimpleNamespace
from datetime import datetime, timedelta
from decimal import Decimal
from typing import Optional

from fastapi import FastAPI, Depends, HTTPException, UploadFile, File, Form, Query, Request, WebSocket, WebSocketDisconnect
from fastapi.exceptions import RequestValidationError
from fastapi.responses import FileResponse, JSONResponse
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from sqlalchemy.orm import Session
from sqlalchemy import func, desc, or_
from sqlalchemy.exc import IntegrityError
from pypdf import PdfReader
from docx import Document
from openpyxl import load_workbook

from . import config, security, schemas, telegram_api, gemini_service, excel_service, location_service, image_markers, payment, google_service as gsvc
from .database import get_db, SessionLocal
from .models import User, Bot, UploadedFile, Message, ExtractedOrder, BotChat, ConversationRead, BotExternalSource

app = FastAPI(title="CuraBot API")
os.makedirs(config.UPLOAD_DIR, exist_ok=True)


def order_is_confirmed(order: dict | None) -> bool:
    return isinstance(order, dict) and order.get("customer_confirmed") is True
app.mount("/uploads", StaticFiles(directory=config.UPLOAD_DIR), name="uploads")
app.add_middleware(
    CORSMiddleware,
    allow_origins=config.CORS_ORIGINS,
    allow_origin_regex=config.CORS_ORIGIN_REGEX or None,
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)


def get_text_from_file(path: str, filename: str) -> str:
    ext = os.path.splitext(filename or "")[1].lower()
    try:
        if ext == ".pdf":
            reader = PdfReader(path)
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        if ext in (".docx", ".doc"):
            doc = Document(path)
            return "\n".join(p.text for p in doc.paragraphs)
        if ext == ".xlsx":
            wb = load_workbook(path, read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if c is None else str(c) for c in row]
                    if any(cells):
                        out.append(" | ".join(cells))
            wb.close()
            return "\n".join(out)
    except Exception:
        return ""
    return ""


def get_bot_or_404(db: Session, user_id: int, bot_id: int) -> Bot:
    bot = db.query(Bot).filter(Bot.id == bot_id, Bot.user_id == user_id, Bot.status == "active").first()
    if not bot:
        raise HTTPException(404, "Bot not found")
    return bot


def get_chat_mode(db: Session, bot_id: int, chat_id: str) -> str:
    row = db.query(BotChat).filter(BotChat.bot_id == bot_id, BotChat.chat_id == str(chat_id)).first()
    return (row.mode if row else "ai")


def set_chat_mode(db: Session, bot_id: int, chat_id: str, mode: str) -> BotChat | None:
    row = db.query(BotChat).filter(BotChat.bot_id == bot_id, BotChat.chat_id == str(chat_id)).first()
    if not row:
        return None
    row.mode = mode
    db.commit()
    db.refresh(row)
    return row


def public_bot(bot: Bot) -> dict:
    return {
        "id": bot.id,
        "name": bot.name,
        "system_prompt": bot.system_prompt,
        "telegram_bot_name": bot.telegram_bot_name,
        "telegram_link": bot.telegram_link,
        "payment_info": bot.payment_info,
        "cash_enabled": bool(bot.cash_enabled),
        "qris_image_url": config.upload_url(bot.id, bot.qris_image_path),
        "status": bot.status,
        "created_at": str(bot.created_at),
    }


def order_to_dict(o: ExtractedOrder) -> dict:
    return {
        "id": o.id,
        "bot_id": o.bot_id,
        "message_id": o.message_id,
        "customer_user_id": o.customer_user_id,
        "customer_name": o.customer_name,
        "products": o.products,
        "total_price": float(o.total_price) if o.total_price is not None else None,
        "delivery_address": o.delivery_address,
        "customer_phone": o.customer_phone,
        "status": o.status,
        "rejection_reason": o.rejection_reason,
        "payment_proof_url": config.upload_url(o.bot_id, o.payment_proof_path),
        "created_at": str(o.created_at),
        "updated_at": str(o.updated_at),
    }


@app.exception_handler(RequestValidationError)
async def validation_handler(request: Request, exc: RequestValidationError):
    errors = {}
    messages = []
    for e in exc.errors():
        field = ".".join(str(p) for p in e.get("loc", []) if p not in ("body", "query", "path"))
        msg = e.get("msg", "Invalid value")
        errors.setdefault(field or "request", []).append(msg)
        messages.append(f"The {field or 'request'} field: {msg}" if field else msg)
    first = messages[0] if messages else "Validation failed."
    summary = first if len(messages) == 1 else f"{first} (and {len(messages) - 1} more error{'s' if len(messages) > 2 else ''})"
    return JSONResponse(status_code=422, content={"message": summary, "errors": errors})


# ---------- AUTH ----------

@app.post("/api/auth/register", status_code=201)
def register(body: schemas.RegisterIn, db: Session = Depends(get_db)):
    if db.query(User).filter(User.email == body.email).first():
        raise HTTPException(409, "Email already registered")
    user = User(email=body.email, password_hash=security.hash_password(body.password))
    db.add(user)
    db.commit()
    db.refresh(user)
    return {"id": user.id, "email": user.email}


@app.post("/api/auth/login")
def login(body: schemas.LoginIn, db: Session = Depends(get_db)):
    user = db.query(User).filter(User.email == body.email).first()
    if not user or not security.verify_password(body.password, user.password_hash):
        raise HTTPException(401, "Invalid email or password")
    return {
        "access_token": security.create_token(user.id),
        "token_type": "bearer",
        "expires_in": config.JWT_EXPIRE_DAYS * 86400,
    }


@app.post("/api/auth/logout")
def logout(user_id: int = Depends(security.get_current_user)):
    return {"ok": True}


# ---------- BOTS ----------

@app.post("/api/bots/create", status_code=201)
def create_bot(body: schemas.BotCreateIn, user_id: int = Depends(security.get_current_user),
               db: Session = Depends(get_db)):
    bot = Bot(
        user_id=user_id, name=body.name, system_prompt=body.system_prompt,
        api_key_encrypted=security.encrypt(body.api_key),
        payment_info=body.payment_info,
        cash_enabled=bool(body.cash_enabled),
        telegram_bot_name=config.TELEGRAM_BOT_USERNAME,
        telegram_link="",
    )
    db.add(bot)
    db.commit()
    db.refresh(bot)
    bot.telegram_link = f"https://t.me/{config.TELEGRAM_BOT_USERNAME}?start=bot_{bot.id}"
    db.commit()
    db.refresh(bot)
    return public_bot(bot)


@app.get("/api/bots/list")
def list_bots(user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    bots = (db.query(Bot).filter(Bot.user_id == user_id, Bot.status == "active")
            .order_by(desc(Bot.created_at)).all())
    return {"bots": [public_bot(b) for b in bots]}


@app.get("/api/bots/{bot_id}")
def get_bot_endpoint(bot_id: int, user_id: int = Depends(security.get_current_user),
                     db: Session = Depends(get_db)):
    return public_bot(get_bot_or_404(db, user_id, bot_id))


@app.put("/api/bots/{bot_id}")
def update_bot(bot_id: int, body: schemas.BotUpdateIn, user_id: int = Depends(security.get_current_user),
               db: Session = Depends(get_db)):
    bot = get_bot_or_404(db, user_id, bot_id)
    if body.name is not None:
        bot.name = body.name
    if body.system_prompt is not None:
        bot.system_prompt = body.system_prompt
    if body.api_key is not None:
        bot.api_key_encrypted = security.encrypt(body.api_key)
    if body.payment_info is not None:
        bot.payment_info = body.payment_info
    if body.cash_enabled is not None:
        bot.cash_enabled = body.cash_enabled
    db.commit()
    db.refresh(bot)
    return public_bot(bot)


@app.delete("/api/bots/{bot_id}")
def delete_bot(bot_id: int, user_id: int = Depends(security.get_current_user),
               db: Session = Depends(get_db)):
    bot = get_bot_or_404(db, user_id, bot_id)
    bot.status = "inactive"
    db.commit()
    return {"ok": True}


# ---------- FILES ----------

@app.post("/api/files/upload", status_code=201)
def upload_file(bot_id: int = Form(...), file: UploadFile = File(...),
                label: str = Form(default=""),
                user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    count = db.query(func.count(UploadedFile.id)).filter(UploadedFile.bot_id == bot_id).scalar()
    if count >= config.MAX_FILES_PER_BOT + config.MAX_IMAGE_FILES_PER_BOT:
        raise HTTPException(400, "File limit reached")
    ext = os.path.splitext(file.filename or "")[1].lower()
    is_image = ext in config.IMAGE_EXTS
    allowed = config.IMAGE_EXTS + (".pdf", ".docx", ".doc", ".xlsx")
    if ext not in allowed:
        raise HTTPException(400, f"Only {', '.join(allowed)} allowed")
    if is_image:
        img_count = (db.query(func.count(UploadedFile.id))
                     .filter(UploadedFile.bot_id == bot_id, UploadedFile.file_type.in_(config.IMAGE_EXTS)).scalar())
        if img_count >= config.MAX_IMAGE_FILES_PER_BOT:
            raise HTTPException(400, f"Max {config.MAX_IMAGE_FILES_PER_BOT} images per bot")
    else:
        doc_count = (db.query(func.count(UploadedFile.id))
                     .filter(UploadedFile.bot_id == bot_id, ~UploadedFile.file_type.in_(config.IMAGE_EXTS)).scalar())
        if doc_count >= config.MAX_FILES_PER_BOT:
            raise HTTPException(400, f"Max {config.MAX_FILES_PER_BOT} files per bot")
    data = file.file.read()
    if len(data) > config.MAX_FILE_SIZE:
        raise HTTPException(400, "File too large (max 25MB)")
    dest_dir = os.path.join(config.UPLOAD_DIR, str(bot_id))
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, f"{int(time.time())}_{os.path.basename(file.filename)}")
    with open(dest, "wb") as f:
        f.write(data)
    text = "" if is_image else get_text_from_file(dest, file.filename)
    row = UploadedFile(bot_id=bot_id, filename=file.filename, file_path=dest,
                       file_type=ext, file_size=len(data), extracted_text=text,
                       label=(label or "").strip() or file.filename)
    db.add(row)
    db.commit()
    db.refresh(row)
    return {"id": row.id, "filename": row.filename, "file_size": row.file_size, "extracted": bool(text)}


@app.post("/api/bots/{bot_id}/qris", status_code=201)
def upload_qris(bot_id: int, file: UploadFile = File(...),
                user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    bot = get_bot_or_404(db, user_id, bot_id)
    ext = os.path.splitext(file.filename or "")[1].lower()
    if ext not in (".jpg", ".jpeg", ".png"):
        raise HTTPException(400, "Only JPG/PNG allowed")
    data = file.file.read()
    if len(data) > config.MAX_FILE_SIZE:
        raise HTTPException(400, "File too large")
    dest_dir = os.path.join(config.UPLOAD_DIR, str(bot_id))
    os.makedirs(dest_dir, exist_ok=True)
    dest = os.path.join(dest_dir, f"qris_{int(time.time())}{ext}")
    with open(dest, "wb") as f:
        f.write(data)
    bot.qris_image_path = dest
    db.commit()
    return {"ok": True, "qris_image_url": config.upload_url(bot_id, dest)}


@app.get("/api/media/{bot_id}/{filename}")
def get_media(bot_id: int, filename: str,
              user_id: int = Depends(security.get_current_user),
              db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    bot_dir = os.path.realpath(os.path.join(config.UPLOAD_DIR, str(bot_id)))
    path = os.path.realpath(os.path.join(bot_dir, os.path.basename(filename)))
    if os.path.dirname(path) != bot_dir or not os.path.isfile(path):
        raise HTTPException(404, "Media not found")
    return FileResponse(path)


@app.get("/api/files/{bot_id}")
def list_files(bot_id: int, user_id: int = Depends(security.get_current_user),
               db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    rows = (db.query(UploadedFile).filter(UploadedFile.bot_id == bot_id)
            .order_by(desc(UploadedFile.created_at)).all())
    return {"files": [{
        "id": r.id, "bot_id": r.bot_id, "filename": r.filename, "file_type": r.file_type,
        "file_size": r.file_size, "created_at": str(r.created_at), "label": r.label,
        "media_url": config.upload_url(r.bot_id, r.file_path),
    } for r in rows]}


@app.post("/api/bots/bulk-delete")
def bulk_delete_bots(payload: schemas.BulkIdsIn, user_id: int = Depends(security.get_current_user),
                     db: Session = Depends(get_db)):
    rows = db.query(Bot).filter(Bot.id.in_(payload.ids), Bot.user_id == user_id).all()
    for bot in rows:
        bot.status = "inactive"
    db.commit()
    return {"deleted": len(rows), "failed": len(payload.ids) - len(rows)}


@app.delete("/api/files/{file_id}")
def delete_file(file_id: int, user_id: int = Depends(security.get_current_user),
                db: Session = Depends(get_db)):
    row = (db.query(UploadedFile).join(Bot, Bot.id == UploadedFile.bot_id)
           .filter(UploadedFile.id == file_id, Bot.user_id == user_id).first())
    if not row:
        raise HTTPException(404, "File not found")
    try:
        resolved = config.resolve_upload_path(row.file_path)
        if resolved:
            os.remove(resolved)
    except OSError:
        pass
    db.delete(row)
    db.commit()
    return {"ok": True}


@app.post("/api/files/bulk-delete")
def bulk_delete_files(payload: schemas.BulkIdsIn, user_id: int = Depends(security.get_current_user),
                      db: Session = Depends(get_db)):
    rows = (db.query(UploadedFile).join(Bot, Bot.id == UploadedFile.bot_id)
            .filter(UploadedFile.id.in_(payload.ids), Bot.user_id == user_id).all())
    for row in rows:
        try:
            resolved = config.resolve_upload_path(row.file_path)
            if resolved:
                os.remove(resolved)
        except OSError:
            pass
        db.delete(row)
    db.commit()
    return {"deleted": len(rows), "failed": len(payload.ids) - len(rows)}


@app.patch("/api/files/{file_id}")
def update_file_label(file_id: int, payload: schemas.FileLabelIn,
                      user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    row = (db.query(UploadedFile).join(Bot, Bot.id == UploadedFile.bot_id)
           .filter(UploadedFile.id == file_id, Bot.user_id == user_id).first())
    if not row:
        raise HTTPException(404, "File not found")
    row.label = payload.label.strip()
    db.commit()
    return {"ok": True, "label": row.label}


# ---------- GOOGLE SOURCES ----------

def source_to_dict(s: BotExternalSource) -> dict:
    return {"id": s.id, "bot_id": s.bot_id, "kind": s.kind, "url": s.url,
            "external_id": s.external_id, "tab": s.tab, "mapping": s.mapping,
            "last_error": s.last_error, "created_at": str(s.created_at)}


@app.get("/api/bots/{bot_id}/google-identity")
def google_identity(bot_id: int, user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    return {"client_email": config.GOOGLE_CLIENT_EMAIL, "configured": bool(config.GOOGLE_CREDENTIALS_PATH)}


@app.post("/api/bots/{bot_id}/sources", status_code=201)
def create_source(bot_id: int, body: schemas.SourceCreateIn, user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    try:
        ext_id = gsvc.parse_google_id(body.kind, body.url)
    except ValueError as e:
        raise HTTPException(400, str(e))
    dup = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot_id, BotExternalSource.kind == body.kind, BotExternalSource.external_id == ext_id).first()
    if dup:
        raise HTTPException(409, "Link ini sudah terhubung")
    row = BotExternalSource(bot_id=bot_id, kind=body.kind, url=body.url.strip()[:2000], external_id=ext_id, tab=(body.tab or "").strip() or None)
    db.add(row)
    db.commit()
    db.refresh(row)
    try:
        if body.kind == "sheet":
            _h, _r, m = gsvc.read_sheet_rows(row)
            row.mapping = {k: v for k, v in m.items() if v is not None}
        else:
            gsvc.list_drive_images(row)
        row.last_error = None
    except Exception as e:
        # ponytail: verifikasi gagal -> hapus row yatim biar retry tidak mentok 409
        err = gsvc._safe_err(e)
        db.delete(row)
        db.commit()
        raise HTTPException(502, f"Verifikasi gagal: {err}. Pastikan link benar & sudah di-share ke {config.GOOGLE_CLIENT_EMAIL}")
    db.commit()
    db.refresh(row)
    return source_to_dict(row)


@app.get("/api/bots/{bot_id}/sources")
def list_sources(bot_id: int, user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    rows = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot_id).order_by(desc(BotExternalSource.created_at)).all()
    return {"sources": [source_to_dict(r) for r in rows], "client_email": config.GOOGLE_CLIENT_EMAIL}


@app.patch("/api/bots/{bot_id}/sources/{source_id}")
def patch_source(bot_id: int, source_id: int, body: schemas.SourceMappingIn, user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    row = db.query(BotExternalSource).filter(BotExternalSource.id == source_id, BotExternalSource.bot_id == bot_id).first()
    if not row:
        raise HTTPException(404, "Source not found")
    if body.tab is not None:
        row.tab = body.tab.strip() or None
    m = dict(row.mapping or {})
    for k in ("name_col", "price_col", "stock_col", "image_col"):
        v = getattr(body, k)
        if v is not None:
            m[k] = v
    row.mapping = m
    gsvc._sheet_cache.pop(bot_id, None)
    db.commit()
    return source_to_dict(row)


@app.delete("/api/bots/{bot_id}/sources/{source_id}")
def delete_source(bot_id: int, source_id: int, user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    row = db.query(BotExternalSource).filter(BotExternalSource.id == source_id, BotExternalSource.bot_id == bot_id).first()
    if not row:
        raise HTTPException(404, "Source not found")
    db.delete(row)
    db.commit()
    gsvc._sheet_cache.pop(bot_id, None)
    gsvc._drive_cache.pop(source_id, None)
    _SYNC_CACHE.pop(source_id, None)
    return {"ok": True}


# ponytail: hasil sync 60 dtk biar polling web murah (Google dipukul maks 1x/mnt/sumber)
_SYNC_CACHE: dict[int, tuple[float, dict]] = {}
_SYNC_CACHE_TTL = 60


@app.post("/api/bots/{bot_id}/sources/{source_id}/sync")
def sync_source(bot_id: int, source_id: int, auto: int = Query(0), user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    row = db.query(BotExternalSource).filter(BotExternalSource.id == source_id, BotExternalSource.bot_id == bot_id).first()
    if not row:
        raise HTTPException(404, "Source not found")
    if auto:
        hit = _SYNC_CACHE.get(row.id)
        if hit and time.time() - hit[0] < _SYNC_CACHE_TTL:
            return hit[1]
    try:
        if row.kind == "sheet":
            headers, rows, m = gsvc.read_sheet_rows(row)
            row.mapping = {k: v for k, v in m.items() if v is not None}
            row.last_error = None
            gsvc._sheet_cache.pop(bot_id, None)
            db.commit()
            nc = m.get("name_col")
            preview = [r[nc] if nc is not None and nc < len(r) else "" for r in rows[:5]]
            payload = {"ok": True, "headers": headers, "mapping": row.mapping, "preview": preview, "total_rows": len(rows)}
        else:
            items = gsvc.list_drive_images(row)
            row.last_error = None
            db.commit()
            payload = {"ok": True, "preview": [i["name"] for i in items[:5]], "total_rows": len(items)}
        _SYNC_CACHE[row.id] = (time.time(), payload)
        return payload
    except Exception as e:
        row.last_error = gsvc._safe_err(e)
        db.commit()
        raise HTTPException(502, f"Sync gagal: {row.last_error}. Pastikan share ke {config.GOOGLE_CLIENT_EMAIL}")


# ---------- MESSAGES ----------

@app.get("/api/messages/{bot_id}")
def get_messages(bot_id: int, page: int = Query(1, ge=1), limit: int = Query(20, le=100),
                 user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    q = db.query(Message).filter(Message.bot_id == bot_id).order_by(desc(Message.created_at))
    total = q.count()
    rows = q.offset((page - 1) * limit).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "messages": [
        {
            "id": r.id, "bot_id": r.bot_id, "user_id": r.user_id, "chat_id": r.chat_id,
            "message_text": r.message_text, "response_text": r.response_text,
            "extracted_data": r.extracted_data, "response_time": r.response_time,
            "model_used": r.model_used, "created_at": str(r.created_at),
        } for r in rows]}


def split_bubbles(r, bot_id):
    bubbles = []
    if r.message_text and r.message_text != "[admin]":
        bubbles.append({
            "id": f"{r.id}-u", "sender": "user", "text": r.message_text,
            "media_url": config.chat_media_url(bot_id, r.media_path),
            "created_at": str(r.created_at),
        })
    if r.response_text:
        bubbles.append({
            "id": f"{r.id}-b", "sender": r.sender if r.sender == "admin" else "bot",
            "text": r.response_text, "media_url": None,
            "created_at": str(r.created_at),
        })
    return bubbles


class BotEventBroker:
    def __init__(self):
        self._lock = threading.Lock()
        self._subscribers: dict[int, list[queue.Queue]] = {}

    def subscribe(self, bot_id: int) -> queue.Queue:
        q = queue.Queue(maxsize=100)
        with self._lock:
            self._subscribers.setdefault(bot_id, []).append(q)
        return q

    def unsubscribe(self, bot_id: int, q: queue.Queue):
        with self._lock:
            subs = self._subscribers.get(bot_id, [])
            if q in subs:
                subs.remove(q)
            if not subs:
                self._subscribers.pop(bot_id, None)

    def publish(self, bot_id: int, event: str, data: dict):
        payload = json.dumps({"type": event, "data": data}, default=str)
        with self._lock:
            subs = list(self._subscribers.get(bot_id, ()))
        for q in subs:
            try:
                q.put_nowait(payload)
            except queue.Full:
                pass


broker = BotEventBroker()


def conversation_item(db: Session, bot_id: int, user_id_value: str) -> dict | None:
    r = (db.query(Message)
         .filter(Message.bot_id == bot_id, Message.user_id == user_id_value)
         .order_by(desc(Message.created_at), desc(Message.id)).first())
    if not r:
        return None
    order = (db.query(ExtractedOrder)
             .filter(ExtractedOrder.bot_id == bot_id,
                     ExtractedOrder.customer_user_id == r.user_id,
                     ExtractedOrder.customer_name.isnot(None))
             .order_by(desc(ExtractedOrder.created_at)).first())
    read = (db.query(ConversationRead)
            .filter(ConversationRead.bot_id == bot_id,
                    ConversationRead.user_id == r.user_id).first())
    incoming = or_(Message.sender.is_(None), Message.sender != "admin")
    if read:
        unread = (db.query(func.count(Message.id)).filter(
            Message.bot_id == bot_id, Message.user_id == r.user_id,
            Message.created_at > read.last_read_at,
            incoming, Message.message_text != "[admin]").scalar() or 0)
    else:
        unread = (db.query(func.count(Message.id)).filter(
            Message.bot_id == bot_id, Message.user_id == r.user_id,
            incoming, Message.message_text != "[admin]").scalar() or 0)
    last = split_bubbles(r, bot_id)
    last_b = last[-1] if last else {"sender": "user", "text": r.message_text}
    binding = db.query(BotChat).filter(BotChat.bot_id == bot_id, BotChat.chat_id == r.chat_id).first()
    return {
        "user_id": r.user_id,
        "customer_name": order.customer_name if order else None,
        "last_text": last_b.get("text") or "",
        "last_sender": last_b.get("sender") or "user",
        "last_created_at": str(r.created_at),
        "unread_count": unread,
        "mode": binding.mode if binding else "ai",
    }


def emit_conversation(db: Session, bot_id: int, user_id_value: str, msg: Message | None = None):
    item = conversation_item(db, bot_id, user_id_value)
    if not item:
        return
    data = {"item": item}
    if msg is not None:
        data["bubbles"] = split_bubbles(msg, bot_id)
    broker.publish(bot_id, "conversation_updated", data)


# ponytail: label gambar terakhir yg dikirim bot per chat, biar reply "ini apa?" terjawab
_LAST_SENT_IMAGES: dict[str, tuple[float, list]] = {}
# ponytail: message_id foto bot -> label, biar reply ke foto tertentu tepat sasaran
_SENT_PHOTO_MSG: dict[int, str] = {}


def _remember_sent_images(chat_id, labels: list) -> None:
    if not labels:
        return
    try:
        _LAST_SENT_IMAGES[str(chat_id)] = (time.time(), [str(label) for label in labels][-5:])
    except Exception:
        pass


def _remember_photo_msg(message_id, label) -> None:
    try:
        if message_id and label:
            _SENT_PHOTO_MSG[int(message_id)] = str(label)
            while len(_SENT_PHOTO_MSG) > 500:
                _SENT_PHOTO_MSG.pop(next(iter(_SENT_PHOTO_MSG)))
    except Exception:
        pass


def _reply_context(message: dict) -> str:
    # ponytail: customer quote-reply -> beri tahu AI 1 baris pesan mana yg dibalas
    try:
        reply = message.get("reply_to_message") or {}
        if not reply:
            return ""
        chat_id = str((message.get("chat") or {}).get("id") or "")
        quoted = ((reply.get("text") or reply.get("caption") or "").strip())[:300]
        if reply.get("photo"):
            try:
                lbl = _SENT_PHOTO_MSG.get(int(reply.get("message_id") or 0)) or ""
            except (TypeError, ValueError):
                lbl = ""
            if not lbl:
                last = _LAST_SENT_IMAGES.get(chat_id)
                lbl = last[1][-1] if last and time.time() - last[0] < 3600 and last[1] else ""
            if lbl:
                ctx = (f"[Customer membalas foto bot (file '{lbl}'). "
                       "Jawab dari nama file ini; bila customer menebak nama lain, cocokkan dulu — "
                       "benarkan hanya bila cocok, koreksi bila tidak. Jangan asal setuju.]")
                return ctx + (f" Kutipan: '{quoted}'." if quoted else "")
            return "[Customer membalas foto dari bot.]" + (f" Kutipan: '{quoted}'." if quoted else "")
        if quoted:
            return f"[Customer membalas pesan: '{quoted}']"
    except Exception:
        pass
    return ""


@app.websocket("/api/bots/{bot_id}/ws")
async def bot_events_ws(websocket: WebSocket, bot_id: int, token: str = Query("")):
    # ponytail: token via query karena WebSocket browser tak bisa set header Authorization
    try:
        user_id = security.decode_token(token)
    except HTTPException:
        await websocket.close(code=4401)
        return
    db = SessionLocal()
    try:
        bot = db.query(Bot).filter(Bot.id == bot_id, Bot.user_id == user_id, Bot.status == "active").first()
    finally:
        db.close()
    if not bot:
        await websocket.close(code=4404)
        return
    q = broker.subscribe(bot_id)
    await websocket.accept()
    try:
        while True:
            payload = await asyncio.to_thread(q.get)
            await websocket.send_text(payload)
    except (WebSocketDisconnect, RuntimeError):
        pass
    finally:
        broker.unsubscribe(bot_id, q)


@app.get("/api/bots/{bot_id}/conversations")
def list_conversations(bot_id: int, q: str = Query("", max_length=100),
                       page: int = Query(1, ge=1), limit: int = Query(20, le=100),
                       user_id: int = Depends(security.get_current_user),
                       db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    sub = (db.query(Message.user_id, func.max(Message.created_at).label("last_at"))
           .filter(Message.bot_id == bot_id).group_by(Message.user_id)
           .subquery())
    rows = (db.query(Message).join(sub,
            (Message.user_id == sub.c.user_id) & (Message.created_at == sub.c.last_at))
            .filter(Message.bot_id == bot_id).order_by(desc(Message.created_at), desc(Message.id)).all())
    seen = set()
    unique_rows = []
    for r in rows:
        if r.user_id in seen:
            continue
        seen.add(r.user_id)
        unique_rows.append(r)
    items = []
    for r in unique_rows:
        order = (db.query(ExtractedOrder)
                 .filter(ExtractedOrder.bot_id == bot_id,
                         ExtractedOrder.customer_user_id == r.user_id,
                         ExtractedOrder.customer_name.isnot(None))
                 .order_by(desc(ExtractedOrder.created_at)).first())
        read = (db.query(ConversationRead)
                .filter(ConversationRead.bot_id == bot_id,
                        ConversationRead.user_id == r.user_id).first())
        unread = 0
        incoming = or_(Message.sender.is_(None), Message.sender != "admin")
        if read:
            unread = (db.query(func.count(Message.id)).filter(
                Message.bot_id == bot_id, Message.user_id == r.user_id,
                Message.created_at > read.last_read_at,
                incoming, Message.message_text != "[admin]").scalar() or 0)
        else:
            unread = (db.query(func.count(Message.id)).filter(
                Message.bot_id == bot_id, Message.user_id == r.user_id,
                incoming, Message.message_text != "[admin]").scalar() or 0)
        last = split_bubbles(r, bot_id)
        last_b = last[-1] if last else {"sender": "user", "text": r.message_text}
        items.append({
            "user_id": r.user_id,
            "customer_name": order.customer_name if order else None,
            "last_text": last_b.get("text") or "",
            "last_sender": last_b.get("sender") or "user",
            "last_created_at": str(r.created_at),
            "unread_count": unread,
        })
    needle = (q or "").strip().lower()
    if needle:
        items = [i for i in items
                 if needle in (i["customer_name"] or "").lower()
                 or needle in str(i["user_id"]).lower()
                 or needle in (i["last_text"] or "").lower()]
    total = len(items)
    start = (page - 1) * limit
    return {"total": total, "page": page, "limit": limit,
            "conversations": items[start:start + limit]}


@app.get("/api/bots/{bot_id}/conversations/{customer_id}/messages")
def get_thread(bot_id: int, customer_id: str, before_id: int = Query(0, ge=0),
               limit: int = Query(30, le=100),
               user_id: int = Depends(security.get_current_user),
               db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    q = db.query(Message).filter(Message.bot_id == bot_id, Message.user_id == customer_id)
    if before_id:
        q = q.filter(Message.id < before_id)
    rows = q.order_by(desc(Message.created_at), desc(Message.id)).limit(limit + 1).all()
    has_more = len(rows) > limit
    rows = rows[:limit]
    bubbles = []
    for r in reversed(rows):
        bubbles.extend(split_bubbles(r, bot_id))
    latest_msg = db.query(Message).filter(Message.bot_id == bot_id, Message.user_id == customer_id).order_by(desc(Message.created_at), desc(Message.id)).first()
    binding = db.query(BotChat).filter(BotChat.bot_id == bot_id, BotChat.chat_id == (latest_msg.chat_id if latest_msg else "")).first() if latest_msg else None
    mode = binding.mode if binding else "ai"
    return {"bubbles": bubbles, "has_more": has_more, "mode": mode}


@app.post("/api/bots/{bot_id}/conversations/{customer_id}/reply")
def reply_thread(bot_id: int, customer_id: str, body: schemas.ReplyIn,
                 user_id: int = Depends(security.get_current_user),
                 db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    latest = (db.query(Message).filter(Message.bot_id == bot_id, Message.user_id == customer_id)
              .order_by(desc(Message.created_at), desc(Message.id)).first())
    if not latest:
        raise HTTPException(404, "Customer chat not found")
    if not body.text.strip():
        raise HTTPException(422, "Reply text empty")
    try:
        telegram_api.send_message(config.TELEGRAM_TOKEN, latest.chat_id, body.text.strip())
    except Exception as e:
        raise HTTPException(502, f"Failed to send Telegram message: {e}")
    msg = Message(bot_id=bot_id, user_id=customer_id, chat_id=latest.chat_id,
                  sender="admin", message_text="[admin]", response_text=body.text.strip())
    db.add(msg)
    db.commit()
    binding = db.query(BotChat).filter(BotChat.bot_id == bot_id, BotChat.chat_id == latest.chat_id).first()
    if binding:
        binding.mode = "manual"
    else:
        db.add(BotChat(bot_id=bot_id, chat_id=latest.chat_id, mode="manual"))
    db.commit()
    emit_conversation(db, bot_id, customer_id, msg)
    return {"ok": True}


@app.post("/api/bots/{bot_id}/conversations/{customer_id}/read")
def mark_read(bot_id: int, customer_id: str,
              user_id: int = Depends(security.get_current_user),
              db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    if not db.query(Message.id).filter(Message.bot_id == bot_id, Message.user_id == customer_id).first():
        raise HTTPException(404, "Customer chat not found")
    row = (db.query(ConversationRead).filter(ConversationRead.bot_id == bot_id,
            ConversationRead.user_id == customer_id).first())
    if not row:
        row = ConversationRead(bot_id=bot_id, user_id=customer_id)
        db.add(row)
        try:
            db.flush()
        except IntegrityError:  # request paralel sudah insert duluan
            db.rollback()
            row = (db.query(ConversationRead).filter(ConversationRead.bot_id == bot_id,
                    ConversationRead.user_id == customer_id).first())
    row.last_read_at = db.query(func.now()).scalar()
    db.commit()
    emit_conversation(db, bot_id, customer_id)
    return {"ok": True}


@app.post("/api/bots/{bot_id}/conversations/{customer_id}/mode")
def set_mode(bot_id: int, customer_id: str, body: schemas.ModeIn,
             user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    latest = (db.query(Message).filter(Message.bot_id == bot_id, Message.user_id == customer_id)
              .order_by(desc(Message.created_at), desc(Message.id)).first())
    if not latest:
        raise HTTPException(404, "Customer chat not found")
    row = set_chat_mode(db, bot_id, latest.chat_id, body.mode)
    if not row:
        raise HTTPException(404, "Customer chat not found")
    emit_conversation(db, bot_id, customer_id, None)
    return {"ok": True, "mode": row.mode}


# ---------- ORDERS ----------

@app.get("/api/bots/{bot_id}/orders")
def get_orders(bot_id: int, page: int = Query(1, ge=1), limit: int = Query(20, le=100),
               status: Optional[str] = None, user_id: int = Depends(security.get_current_user),
               db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    q = db.query(ExtractedOrder).filter(ExtractedOrder.bot_id == bot_id)
    if status:
        q = q.filter(ExtractedOrder.status == status)
    total = q.count()
    rows = q.order_by(desc(ExtractedOrder.created_at)).offset((page - 1) * limit).limit(limit).all()
    return {"page": page, "limit": limit, "total": total, "orders": [order_to_dict(r) for r in rows]}


@app.put("/api/bots/{bot_id}/orders/{order_id}")
def update_order(bot_id: int, order_id: int, body: schemas.OrderStatusIn,
                 user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    order = db.query(ExtractedOrder).filter(ExtractedOrder.id == order_id,
                                            ExtractedOrder.bot_id == bot_id).first()
    if not order:
        raise HTTPException(404, "Order not found")
    prev = order.status
    order.status = body.status
    order.rejection_reason = body.reason if body.status == "rejected" else None
    if body.status == "rejected" and prev != "rejected" and order.stock_deducted:
        try:
            sheets = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot_id, BotExternalSource.kind == "sheet").all()
            for _p in (order.products or []):
                for _s in sheets:
                    try:
                        if gsvc.adjust_stock(_s, str(_p.get("product_name", "")), int(_p.get("quantity", 0) or 0)):
                            break
                    except Exception:
                        continue
            order.stock_deducted = False
        except Exception:
            pass
    db.commit()

    # webhook ke customer: kabari perubahan status order
    msg_row = db.query(Message).filter(Message.id == order.message_id).first() if order.message_id else None
    if msg_row and msg_row.chat_id:
        status_text = {
            "pending": "menunggu konfirmasi penjual",
            "incomplete": "menunggu kelengkapan data pesanan",
            "confirmed": "sudah dikonfirmasi dan sedang diproses",
            "shipped": "sudah dikirim",
            "completed": "sudah selesai",
            "rejected": "ditolak",
        }.get(body.status, body.status)
        note = f"Halo! Status pesanan Anda saat ini: {status_text}."
        if body.status == "rejected":
            note += f"\nAlasan: {body.reason or '-'}\nSilakan hubungi kami jika ingin memesan ulang."
        elif body.status == "pending":
            note += " Mohon tunggu, kami akan mengabari Anda lagi setelah ada update."
        elif body.status == "incomplete":
            missing = []
            if not order.customer_phone:
                missing.append("nomor handphone")
            if not order.delivery_address:
                missing.append("alamat pengiriman")
            if not (order.products or []):
                missing.append("daftar produk")
            detail = ", ".join(missing) if missing else "data pesanan"
            note += f"\nMohon lengkapi data berikut agar pesanan bisa diproses: {detail}."
        try:
            telegram_api.send_message(config.TELEGRAM_TOKEN, msg_row.chat_id, note)
        except Exception:
            pass
    return {"ok": True}


@app.post("/api/bots/{bot_id}/orders/{order_id}/remind")
def remind_payment(bot_id: int, order_id: int, user_id: int = Depends(security.get_current_user),
                   db: Session = Depends(get_db)):
    bot = get_bot_or_404(db, user_id, bot_id)
    order = db.query(ExtractedOrder).filter(ExtractedOrder.id == order_id,
                                            ExtractedOrder.bot_id == bot_id).first()
    if not order:
        raise HTTPException(404, "Order not found")
    if order.payment_proof_path:
        raise HTTPException(400, "Payment proof already received")
    chat_id = None
    if order.message_id:
        msg = db.query(Message).filter(Message.id == order.message_id).first()
        if msg and msg.chat_id:
            chat_id = msg.chat_id
    if not chat_id:
        latest_msg = (db.query(Message)
                      .filter(Message.bot_id == bot_id, Message.user_id == order.customer_user_id)
                      .order_by(desc(Message.created_at)).first())
        if latest_msg and latest_msg.chat_id:
            chat_id = latest_msg.chat_id
    if not chat_id:
        raise HTTPException(404, "Customer chat not found")
    total = f"Rp{float(order.total_price):,.0f}".replace(",", ".") if order.total_price is not None else "-"
    text = (f"Halo {order.customer_name or 'Kak'}! Ini pengingat pembayaran untuk pesanan #{order.id} "
            f"sebesar {total}.\n")
    if bot.payment_info:
        text += f"{bot.payment_info}\n"
    text += "Mohon segera membayar dan kirimkan foto/screenshot bukti pembayaran di chat ini ya. Terima kasih!"
    try:
        telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id, text)
    except Exception as e:
        raise HTTPException(502, f"Failed to send Telegram reminder: {e}")
    return {"ok": True}


@app.get("/api/bots/{bot_id}/orders/export")
def export_orders(bot_id: int, user_id: int = Depends(security.get_current_user),
                  db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    path = os.path.join(config.UPLOAD_DIR, str(bot_id), "orders.xlsx")
    if not os.path.exists(path):
        raise HTTPException(404, "No orders exported yet")
    return FileResponse(path, filename="orders.xlsx",
                        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


# ---------- ANALYTICS ----------

@app.get("/api/bots/{bot_id}/analytics")
def analytics(bot_id: int, range: int = Query(7, ge=1, le=365),
              user_id: int = Depends(security.get_current_user), db: Session = Depends(get_db)):
    get_bot_or_404(db, user_id, bot_id)
    since = datetime.utcnow() - timedelta(days=range)

    total_conv = db.query(func.count(Message.id)).filter(Message.bot_id == bot_id, Message.created_at >= since).scalar()
    unique_users = db.query(func.count(func.distinct(Message.user_id))).filter(Message.bot_id == bot_id, Message.created_at >= since).scalar()
    avg_rt = db.query(func.avg(Message.response_time)).filter(Message.bot_id == bot_id, Message.created_at >= since).scalar()
    total_orders = db.query(func.count(ExtractedOrder.id)).filter(ExtractedOrder.bot_id == bot_id, ExtractedOrder.created_at >= since).scalar()
    orders_extracted = db.query(func.count(ExtractedOrder.id)).filter(ExtractedOrder.bot_id == bot_id, ExtractedOrder.status != "incomplete", ExtractedOrder.created_at >= since).scalar()

    model_rows = (db.query(Message.model_used, func.count(Message.id))
                  .filter(Message.bot_id == bot_id, Message.created_at >= since)
                  .group_by(Message.model_used).all())
    total_m = sum(c for _, c in model_rows) or 1
    success = sum(c for m, c in model_rows if m in ("flash-3.6", "flash-3.5-lite"))

    per_day = (db.query(func.date(Message.created_at), func.count(Message.id))
               .filter(Message.bot_id == bot_id, Message.created_at >= since)
               .group_by(func.date(Message.created_at)).order_by(func.date(Message.created_at)).all())
    orders_day = (db.query(func.date(ExtractedOrder.created_at), func.count(ExtractedOrder.id))
                  .filter(ExtractedOrder.bot_id == bot_id, ExtractedOrder.created_at >= since)
                  .group_by(func.date(ExtractedOrder.created_at)).order_by(func.date(ExtractedOrder.created_at)).all())

    return {
        "total_conversations": total_conv,
        "unique_users": unique_users,
        "avg_response_time": round(float(avg_rt), 2) if avg_rt else 0,
        "total_orders": total_orders,
        "orders_extracted": orders_extracted,
        "model_success_rate": round(success / total_m * 100, 1),
        "conversations_per_day": [{"date": str(d), "count": c} for d, c in per_day],
        "orders_per_day": [{"date": str(d), "count": c} for d, c in orders_day],
    }


# ---------- TELEGRAM WEBHOOK ----------

@app.post("/api/telegram/webhook")
async def telegram_webhook(request: Request, db: Session = Depends(get_db)):
    if config.TELEGRAM_WEBHOOK_SECRET:
        provided = request.headers.get("X-Telegram-Bot-Api-Secret-Token", "")
        if not hmac.compare_digest(provided, config.TELEGRAM_WEBHOOK_SECRET):
            raise HTTPException(403, "Invalid webhook secret")
    payload = await request.json()
    # tap pada tombol metode pembayaran -> jawab deterministik tanpa AI
    cb = payload.get("callback_query") or {}
    cb_data = cb.get("data") or ""
    if cb_data.startswith("pay:"):
        cb_chat_id = cb.get("message", {}).get("chat", {}).get("id")
        try:
            telegram_api.answer_callback_query(config.TELEGRAM_TOKEN, cb.get("id"))
        except Exception:
            pass
        binding = db.query(BotChat).filter(BotChat.chat_id == str(cb_chat_id)).first()
        bot = (db.query(Bot).filter(Bot.id == binding.bot_id, Bot.status == "active").first()
               if binding else None)
        if bot:
            kind, text = payment.callback_response(bot, cb_data)
            try:
                if kind == "qris":
                    telegram_api.send_photo(config.TELEGRAM_TOKEN, cb_chat_id,
                                            config.resolve_upload_path(bot.qris_image_path),
                                            caption="QRIS pembayaran")
                telegram_api.send_message(config.TELEGRAM_TOKEN, cb_chat_id, text)
                msg = Message(bot_id=bot.id, user_id=str(cb.get("from", {}).get("id") or cb_chat_id), chat_id=str(cb_chat_id),
                              sender="user",
                              message_text=f"[memilih metode pembayaran: {cb_data.removeprefix('pay:')}]",
                              response_text=text)
                db.add(msg)
                db.commit()
                db.refresh(msg)
                emit_conversation(db, bot.id, str(cb_chat_id), msg)
            except Exception:
                pass
        return {"ok": True}
    message = payload.get("message") or {}
    chat_id = message.get("chat", {}).get("id")
    # customer share contact -> isi otomatis customer_phone di order aktif/terakhir
    contact = message.get("contact")
    if contact and contact.get("phone_number"):
        binding = db.query(BotChat).filter(BotChat.chat_id == str(chat_id)).first()
        if binding:
            latest = (db.query(ExtractedOrder)
                      .join(Message, Message.id == ExtractedOrder.message_id)
                      .filter(ExtractedOrder.bot_id == binding.bot_id, Message.chat_id == str(chat_id))
                      .order_by(desc(ExtractedOrder.created_at)).first())
            if latest and not latest.customer_phone:
                latest.customer_phone = contact["phone_number"]
                db.commit()
        return {"ok": True}
    text = (message.get("text") or "").strip()
    # ponytail: reply/quote -> tempel konteks yg dibalas biar AI nyambung ("ini apa?", "mana?")
    reply_ctx = _reply_context(message)
    if reply_ctx and text:
        text = f"{reply_ctx} {text}"
    # customer mengirim foto bukti pembayaran -> simpan & teruskan ke dashboard, status TIDAK diubah AI
    photos = message.get("photo") or []
    if photos:
        caption = (message.get("caption") or "").strip()
        if reply_ctx:
            caption = f"{reply_ctx} {caption}".strip()
        chat_id = message.get("chat", {}).get("id")
        binding = db.query(BotChat).filter(BotChat.chat_id == str(chat_id)).first()
        if binding:
            bot = db.query(Bot).filter(Bot.id == binding.bot_id, Bot.status == "active").first()
            if bot:
                if (binding.mode or "ai") == "manual":
                    try:
                        fid = photos[-1]["file_id"]
                        info = telegram_api.get_file(config.TELEGRAM_TOKEN, fid)
                        data = telegram_api.download_file(config.TELEGRAM_TOKEN, info["file_path"])
                        dest_dir = os.path.join(config.UPLOAD_DIR, str(bot.id))
                        os.makedirs(dest_dir, exist_ok=True)
                        dest = os.path.join(dest_dir, f"proof_{int(time.time())}.jpg")
                        with open(dest, "wb") as f:
                            f.write(data)
                        photo_msg = Message(bot_id=bot.id, user_id=str(chat_id), chat_id=str(chat_id),
                                           sender="user", media_path=dest,
                                           message_text=caption or "[foto]", response_text="")
                        db.add(photo_msg)
                        db.commit()
                        db.refresh(photo_msg)
                        emit_conversation(db, bot.id, str(chat_id), photo_msg)
                    except Exception:
                        telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id,
                                                  "Maaf, bukti tidak bisa kami proses. Coba kirim ulang ya.")
                    return {"ok": True}
                try:
                    fid = photos[-1]["file_id"]
                    info = telegram_api.get_file(config.TELEGRAM_TOKEN, fid)
                    data = telegram_api.download_file(config.TELEGRAM_TOKEN, info["file_path"])
                    dest_dir = os.path.join(config.UPLOAD_DIR, str(bot.id))
                    os.makedirs(dest_dir, exist_ok=True)
                    dest = os.path.join(dest_dir, f"proof_{int(time.time())}.jpg")
                    with open(dest, "wb") as f:
                        f.write(data)
                    photo_msg = Message(bot_id=bot.id, user_id=str(chat_id), chat_id=str(chat_id),
                                       sender="user", media_path=dest,
                                       message_text=caption or "[foto]",
                                       response_text="")
                    db.add(photo_msg)
                    db.commit()
                    db.refresh(photo_msg)
                    emit_conversation(db, bot.id, str(chat_id), photo_msg)
                    api_key = security.decrypt(bot.api_key_encrypted)
                    files = (db.query(UploadedFile)
                             .filter(UploadedFile.bot_id == bot.id, UploadedFile.extracted_text != "").all())
                    kb = "\n\n".join((f.extracted_text or "")[:1500] for f in files)
                    try:
                        is_proof, answer = gemini_service.analyze_photo(api_key, data, caption, kb)
                    except Exception:
                        is_proof, answer = True, None  # ponytail: gagal cek -> jangan blokir pelanggan, anggap bukti
                    if not is_proof:
                        if answer:
                            try:
                                telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id, answer)
                            except Exception:
                                pass
                            photo_msg.response_text = answer
                            db.commit()
                            emit_conversation(db, bot.id, str(chat_id), photo_msg)
                            return {"ok": True}
                        if caption:
                            # ponytail: caption ada tapi analyze_photo tak menjawab -> kirim ke vision lagi dengan gambar
                            try:
                                files2 = files
                                rows = (db.query(Message)
                                        .filter(Message.bot_id == bot.id, Message.chat_id == str(chat_id),
                                                Message.id < photo_msg.id)
                                        .order_by(desc(Message.created_at)).limit(4).all())
                                hist = []
                                for r in reversed(rows):
                                    hist.append({"role": "user", "text": r.message_text})
                                    if r.response_text:
                                        hist.append({"role": "model", "text": r.response_text})
                                reply2, _m = gemini_service.answer_with_image(
                                    api_key, bot.system_prompt or "", kb, hist, data, caption)
                                telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id, reply2)
                                photo_msg.response_text = reply2
                                db.commit()
                                emit_conversation(db, bot.id, str(chat_id), photo_msg)
                                return {"ok": True}
                            except Exception:
                                text = caption  # fallback terakhir: alur teks
                        else:
                            try:
                                telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id, gemini_service.PHOTO_ASK_BACK)
                            except Exception:
                                pass
                            photo_msg.response_text = gemini_service.PHOTO_ASK_BACK
                            db.commit()
                            emit_conversation(db, bot.id, str(chat_id), photo_msg)
                            return {"ok": True}
                    else:
                        latest = (db.query(ExtractedOrder)
                                  .join(Message, Message.id == ExtractedOrder.message_id)
                                  .filter(ExtractedOrder.bot_id == bot.id, Message.chat_id == str(chat_id),
                                          ExtractedOrder.status.in_(["pending", "incomplete"]))
                                  .order_by(desc(ExtractedOrder.created_at)).first())
                        if latest:
                            latest.payment_proof_path = dest
                            db.commit()
                        telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id,
                                                  "Terima kasih, bukti pembayaran sudah kami terima dan diteruskan ke penjual. Mohon tunggu konfirmasi ya.")
                        if not caption:
                            return {"ok": True}
                        try:
                            rows = (db.query(Message)
                                    .filter(Message.bot_id == bot.id, Message.chat_id == str(chat_id),
                                            Message.id < photo_msg.id)
                                    .order_by(desc(Message.created_at)).limit(4).all())
                            hist = []
                            for r in reversed(rows):
                                hist.append({"role": "user", "text": r.message_text})
                                if r.response_text:
                                    hist.append({"role": "model", "text": r.response_text})
                            reply2, _m = gemini_service.answer_with_image(
                                api_key, bot.system_prompt or "", kb, hist, data, caption)
                            telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id, reply2)
                            photo_msg.response_text = reply2
                            db.commit()
                            emit_conversation(db, bot.id, str(chat_id), photo_msg)
                            return {"ok": True}
                        except Exception:
                            text = caption
                except Exception:
                    telegram_api.send_message(config.TELEGRAM_TOKEN, chat_id,
                                              "Maaf, bukti tidak bisa kami proses. Coba kirim ulang ya.")
        if not caption:
            return {"ok": True}
        text = caption
    if not text:
        return {"ok": True}
    chat_id = message.get("chat", {}).get("id")
    from_user = message.get("from", {})
    uid = str(from_user.get("id", ""))
    customer_name = " ".join(filter(None, [from_user.get("first_name"), from_user.get("last_name")])) or None
    token = config.TELEGRAM_TOKEN

    if text.lower().startswith("/start"):
        arg = text.split(maxsplit=1)[1] if len(text.split(maxsplit=1)) > 1 else ""
        if arg.startswith("bot_"):
            try:
                target_id = int(arg[4:])
            except ValueError:
                target_id = None
            bot = (db.query(Bot).filter(Bot.id == target_id, Bot.status == "active").first()
                   if target_id else None)
            if bot:
                existing = db.query(BotChat).filter(BotChat.chat_id == str(chat_id)).first()
                if existing:
                    existing.bot_id = bot.id
                else:
                    db.add(BotChat(bot_id=bot.id, chat_id=str(chat_id)))
                db.commit()
                try:
                    telegram_api.send_message(token, chat_id, f"Connected to {bot.name}. How can I help you?")
                except Exception:
                    pass
                return {"ok": True}
        try:
            telegram_api.send_message(token, chat_id, "Welcome! Please use a bot link shared by the business to start.")
        except Exception:
            pass
        return {"ok": True}

    binding = db.query(BotChat).filter(BotChat.chat_id == str(chat_id)).first()
    if not binding:
        try:
            telegram_api.send_message(token, chat_id, "Please start via a bot link shared by the business.")
        except Exception:
            pass
        return {"ok": True}
    bot = db.query(Bot).filter(Bot.id == binding.bot_id, Bot.status == "active").first()
    if not bot:
        return {"ok": True}
    bot_id = bot.id
    if (binding.mode or "ai") == "manual":
        msg = Message(bot_id=bot_id, user_id=uid, chat_id=str(chat_id),
                      sender="user", message_text=text, response_text="")
        db.add(msg)
        db.commit()
        db.refresh(msg)
        emit_conversation(db, bot_id, uid, msg)
        return {"ok": True}
    api_key = security.decrypt(bot.api_key_encrypted)

    # foto sebelumnya ditanya balik? -> jawaban teks user ini merujuk ke foto itu
    pending_photo = (db.query(Message)
                     .filter(Message.bot_id == bot_id, Message.chat_id == str(chat_id),
                             Message.media_path.isnot(None),
                             Message.response_text == gemini_service.PHOTO_ASK_BACK)
                     .order_by(desc(Message.id)).first())
    if pending_photo:
        answered = (db.query(Message)
                    .filter(Message.bot_id == bot_id, Message.chat_id == str(chat_id),
                            Message.id > pending_photo.id, Message.media_path.is_(None))
                    .first())
        if answered:
            pending_photo = None

    history_rows = (db.query(Message)
                    .filter(Message.bot_id == bot_id, Message.chat_id == str(chat_id))
                    .order_by(desc(Message.created_at)).limit(5).all())
    history = []
    for r in reversed(history_rows):
        history.append({"role": "user", "text": r.message_text})
        history.append({"role": "model", "text": r.response_text})

    files = (db.query(UploadedFile)
             .filter(UploadedFile.bot_id == bot_id, UploadedFile.extracted_text != "").all())
    file_context = "\n\n".join((f.extracted_text or "")[:3000] for f in files)
    try:
        gctx = gsvc.sheet_context(bot.id, db)
        if gctx:
            file_context = (file_context + "\n\n" + gctx)[:15000]
    except Exception:
        pass

    # status order terakhir untuk konteks jawaban AI
    latest_order = (db.query(ExtractedOrder)
                    .join(Message, Message.id == ExtractedOrder.message_id)
                    .filter(ExtractedOrder.bot_id == bot_id, Message.chat_id == str(chat_id))
                    .order_by(desc(ExtractedOrder.created_at)).first())
    system_prompt = bot.system_prompt
    if latest_order:
        guidance = {
            "pending": "katakan pesanan sedang menunggu konfirmasi penjual, minta customer menunggu, akan diinformasikan lagi.",
            "incomplete": "katakan data pesanan belum lengkap, minta customer melengkapi data yang kurang saja (lihat detail di bawah).",
            "confirmed": "katakan pesanan sudah dikonfirmasi dan sedang diproses.",
            "shipped": "katakan pesanan sudah dikirim.",
            "completed": "katakan pesanan sudah selesai, tawarkan belanja lagi.",
            "rejected": f"katakan pesanan ditolak. Alasan: {latest_order.rejection_reason or '-'}. Tawarkan pesan ulang.",
        }.get(latest_order.status, "")
        if latest_order.status == "incomplete":
            missing = []
            if not latest_order.customer_phone:
                missing.append("nomor handphone")
            # alamat hanya ditanya jika layanan delivery tersedia dan alamat belum diisi; jika "Onsite", jangan tanya
            if not latest_order.delivery_address:
                missing.append("alamat pengiriman (hanya untuk pengiriman/delivery; jika pesan di tempat, isi 'Onsite')")
            if not (latest_order.products or []):
                missing.append("daftar produk")
            if latest_order.delivery_address and latest_order.delivery_address.strip().lower() == "onsite":
                missing = [m for m in missing if not m.startswith("alamat")]
            guidance += f" Data yang masih kurang: {', '.join(missing) if missing else 'tidak ada — data sudah lengkap'}."
        system_prompt += f"\n\nINFO STATUS PESANAN CUSTOMER SAAT INI: status='{latest_order.status}'. Jika customer bertanya status pesanan, {guidance}"
        detail = {"products": latest_order.products, "total_price": float(latest_order.total_price) if latest_order.total_price else None,
                  "delivery_address": latest_order.delivery_address, "customer_phone": latest_order.customer_phone}
        system_prompt += (f"\n\nDETAIL PESANAN CUSTOMER: {detail}. Jika customer bertanya apa saja pesanannya, "
                          "makan di tempat atau tidak, harga, atau data pesanan lainnya, JAWAB LANGSUNG dari data ini, jangan hanya menyebutkan status.")
        if latest_order.status in ("rejected", "completed", "shipped"):
            system_prompt += ("\n\nPENTING: Jangan mencatat pesanan baru kecuali customer secara eksplisit menyebut produk yang ingin dipesan ulang. "
                              "Kalau customer hanya menanggapi/merespon pesan, balas singkat saja tanpa mencatat pesanan.")
    system_prompt += ("\n\nATURAN KONFIRMASI: Sebelum menyatakan pesanan sudah dicatat/diubah, SELALU ulangi ringkasan lengkap pesanan (produk, jumlah, harga, dan catatan/note per item masing-masing) dan minta konfirmasi customer (\"Benar, Kak?\"). "
                      "Baru setelah customer mengiyakan, katakan pesanan sudah tercatat/terupdate. "
                      "Kalau customer meralat, jawab ringkasan baru dan minta konfirmasi lagi.")
    data_complete = payment.order_data_complete(latest_order)
    if bot.payment_info and data_complete and latest_order.status in ("pending", "incomplete"):
        total_str = f"Rp{float(latest_order.total_price):,.0f}".replace(",", ".")
        system_prompt += ("\n\nTAHAP PEMBAYARAN: Data pesanan customer sudah lengkap. "
                          "Tegaskan dan ingatkan customer agar segera membayar sekarang. "
                          f"Total yang harus dibayar: {total_str}. "
                          "Customer melihat tombol pilihan metode pembayaran di bawah pesan. "
                          "JANGAN uraikan detail tiap metode (nomor rekening, QRIS, cash) dalam teks — "
                          "cukup sebutkan total, minta customer menekan salah satu tombol, dan ingatkan kirim bukti bayar. "
                          "Detail tiap metode hanya dijawab bila customer bertanya eksplisit via teks, dari: " + bot.payment_info)
        if not latest_order.payment_proof_path:
            system_prompt += ("\n\nPENGINGAT BUKTI BAYAR: Customer belum mengirim bukti pembayaran. "
                              "Di setiap balasan, akhiri dengan ajakan tegas mengirimkan foto/screenshot "
                              "bukti pembayaran di chat ini sebelum pesanan diproses. "
                              "Ikuti gaya bahasa pada system prompt.")
    has_payment_info = bool((bot.payment_info or "").strip())
    has_qris = bool(bot.qris_image_path)
    pay_state = f"ada: {bot.payment_info.strip()}" if has_payment_info else "KOSONG (penjual belum mengisi info pembayaran)"
    qris_state = "TERSEDIA (gambar QRIS penjual sudah diupload)" if has_qris else "TIDAK TERSEDIA (penjual belum upload gambar QRIS)"
    cash_state = "TERSEDIA" if bot.cash_enabled else "TIDAK TERSEDIA"
    system_prompt += (
        "\n\nINFO PEMBAYARAN (ground truth, prioritas tertinggi untuk topik bayar): "
        f"payment_info={pay_state}; qris_image={qris_state}; cash={cash_state}."
        "\nATURAN PEMBAYARAN (wajib, kalahkan instruksi lain): "
        "1. Jawab bayar/transfer/rekening/nomor rekening/QRIS HANYA dari info di atas. "
        "2. DILARANG mengarang nomor rekening/bank/VA/QRIS. "
        "3. Tautan toko/website di system prompt BUKAN metode pembayaran — JANGAN sebut/arahkan ke link toko saat ditanya pembayaran, "
        "kecuali payment_info di atas eksplisit memuat link itu sebagai cara bayar. "
        "4. Ditanya rekening: payment_info ada -> kutip persis; kosong -> katakan info pembayaran belum diisi penjual, minta tunggu, jangan arahkan ke website. "
        "5. Ditanya QRIS: TERSEDIA -> katakan QRIS tersedia dan gambarnya menyusul tepat setelah pesan ini; "
        "TIDAK TERSEDIA -> katakan QRIS belum tersedia, tawarkan payment_info bila ada, bila kosong katakan tunggu info penjual. "
        "6. Setiap customer bertanya soal pembayaran/cara bayar (umum, bukan spesifik satu metode): "
        "sistem otomatis menampilkan tombol metode pembayaran di bawah pesanmu — jawab SINGKAT saja "
        "(contoh: 'Gampang, Kak! Silakan pilih metode pembayaran lewat tombol di bawah ya'), "
        "JANGAN uraikan daftar/detail tiap metode dalam teks kecuali customer meminta eksplisit. "
        "Bila semua metode kosong -> katakan info pembayaran belum diisi penjual, minta tunggu. "
        "Jangan pernah katakan 'belum tersedia' untuk hal yang ground truth sebut TERSEDIA/ada. "
        "7. Ditanya bayar tunai/cash/COD: TERSEDIA -> katakan bisa bayar tunai saat pesanan diterima; "
        "TIDAK TERSEDIA -> jangan tawarkan tunai, arahkan ke metode yang ada. "
    )

    image_rows = (db.query(UploadedFile)
                  .filter(UploadedFile.bot_id == bot.id, UploadedFile.file_type.in_(config.IMAGE_EXTS))
                  .order_by(UploadedFile.created_at.asc()).all())
    # ponytail: AI baca Drive otomatis — daftar nama file Drive masuk prompt, tanpa set manual
    drive_names: list[str] = []
    try:
        for ds in db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot.id, BotExternalSource.kind == "drive_folder").all():
            try:
                drive_names.extend(i["name"] for i in gsvc.list_drive_images(ds) if i.get("name"))
            except Exception:
                continue
    except Exception:
        pass
    if image_rows or drive_names:
        lines = "\n".join(f"- {(r.label or r.filename)} | {r.filename}" for r in image_rows)
        if drive_names:
            lines += ("\n" if lines else "") + "\n".join(f"- {n}" for n in drive_names[:100])
        system_prompt += (
            "\n\nDAFTAR GAMBAR TERSEDIA (kirim ke customer hanya bila relevan dan diminta):\n" + lines + "\n"
            "ATURAN GAMBAR (wajib):\n"
            "1. Jika customer meminta melihat gambar produk/menu/kolase, jawab singkat, lalu tambahkan baris `[IMG: label]` dengan label PERSIS seperti di daftar untuk tiap gambar yang cocok. Boleh juga pakai filename bila label kosong.\n"
            "2. DILARANG mengarang label/filename yang tidak ada di daftar.\n"
            "3. Tidak ada gambar yang cocok -> katakan belum tersedia, JANGAN tulis marker.\n"
            "4. Customer minta 'semua gambar' -> tulis marker untuk semua gambar di daftar.\n"
            "5. Gambar tidak ada di daftar -> JANGAN berjanji mengirim foto ('ini fotonya', 'saya kirim ulang'); cukup SATU pesan singkat bahwa gambarnya belum tersedia.\n"
            "6. Customer menebak nama ('ini X bukan?', 'kalau ini?') -> cocokkan X dengan nama file foto/daftar; benarkan hanya bila cocok, koreksi bila tidak. Jangan asal setuju."
        )

    start = time.time()
    try:
        if pending_photo:
            try:
                with open(config.resolve_upload_path(pending_photo.media_path), "rb") as f:
                    img_bytes = f.read()
                reply, model_used = gemini_service.answer_with_image(
                    api_key, system_prompt, file_context, history, img_bytes, text)
            except Exception:
                reply, model_used = gemini_service.chat_with_fallback(api_key, system_prompt, file_context, history, text)
        else:
            reply, model_used = gemini_service.chat_with_fallback(api_key, system_prompt, file_context, history, text)
        m_coords = re.search(r"\(koordinat:\s*([^)]+)\)", bot.system_prompt or "")
        reply = location_service.enrich_with_maps(reply, m_coords.group(1).strip() if m_coords else None)
    except Exception:
        reply, model_used = "Sorry, there is a temporary issue. Please try again in a moment.", "none"
    response_time = round(time.time() - start, 2)

    clean_reply, wanted = image_markers.parse_image_markers(reply)
    # tombol metode: selama tahap pembayaran, atau kapan pun customer menyinggung topik bayar
    show_pay = (payment.available_methods(bot)
                and (payment.payment_stage_active(bot, latest_order)
                     or payment.mentions_payment(text)))
    pay_kb = payment.payment_keyboard(bot) if show_pay else None
    paths = []
    tmp_files = []
    sent_labels = []
    for name in wanted:
        # ponytail: toleransi format lama "label | filename" -> coba bagian filename juga
        candidates = [name] + ([name.rsplit("|", 1)[-1].strip()] if "|" in name else [])
        row = None
        for cand in candidates:
            row = (db.query(UploadedFile)
                   .filter(UploadedFile.bot_id == bot.id,
                           ((UploadedFile.label == cand) | (UploadedFile.filename == cand)),
                           UploadedFile.file_type.in_(config.IMAGE_EXTS)).first())
            if row:
                break
        if row:
            p = config.resolve_upload_path(row.file_path)
            if p and os.path.exists(p):
                paths.append(p)
                sent_labels.append(row.label or row.filename)
    if wanted:
        have = set(os.path.basename(p) for p in paths)
        missing = [w for w in wanted if w not in have]
        if missing:
            try:
                dsources = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot.id, BotExternalSource.kind == "drive_folder").all()
                for ds in dsources:
                    try:
                        items = gsvc.list_drive_images(ds)
                        names = [i["name"] for i in items]
                        for w in list(missing):
                            hit = gsvc.fuzzy_find(w, names)
                            if not hit:
                                continue
                            fid = next(i["id"] for i in items if i["name"] == hit)
                            data = gsvc.download_drive_image(fid)
                            tmp = f"/tmp/curabot-{bot.id}-{int(time.time() * 1000)}-{os.path.basename(hit)}"
                            with open(tmp, "wb") as _f:
                                _f.write(data)
                            # ponytail: JANGAN kirim di sini; kumpulkan lalu kirim sekali di bawah
                            paths.append(tmp)
                            tmp_files.append(tmp)
                            sent_labels.append(hit)
                            missing.remove(w)
                    except Exception:
                        continue
            except Exception:
                pass
    # ponytail: kirim sekali; gagal/tidak-tersedia -> SATU pesan gabungan (anti dobel)
    image_note_sent = False
    if paths:
        delivered = 0
        for i in range(0, len(paths), 10):
            chunk = paths[i:i + 10]
            chunk_labels = sent_labels[i:i + 10]
            try:
                if len(chunk) == 1:
                    res = telegram_api.send_photo(token, chat_id, chunk[0])
                    _remember_photo_msg((res or {}).get("message_id"), chunk_labels[0] if chunk_labels else "")
                else:
                    try:
                        results = telegram_api.send_media_group(token, chat_id, chunk)
                        for r, lb in zip(results or [], chunk_labels):
                            _remember_photo_msg((r or {}).get("message_id"), lb)
                    except Exception:
                        for p, lb in zip(chunk, chunk_labels):
                            res1 = telegram_api.send_photo(token, chat_id, p)
                            _remember_photo_msg((res1 or {}).get("message_id"), lb)
                delivered += len(chunk)
            except Exception:
                for p, lb in zip(chunk, chunk_labels):
                    try:
                        res1 = telegram_api.send_photo(token, chat_id, p)
                        _remember_photo_msg((res1 or {}).get("message_id"), lb)
                        delivered += 1
                    except Exception:
                        continue
        for t in tmp_files:
            try:
                os.remove(t)
            except OSError:
                pass
        if delivered:
            _remember_sent_images(chat_id, sent_labels or wanted)
        else:
            note = "Maaf, gambar gagal dikirim. Coba lagi ya."
            try:
                telegram_api.send_message(token, chat_id,
                                          f"{clean_reply}\n\n{note}" if clean_reply else note,
                                          reply_markup=pay_kb)
                image_note_sent = True
            except Exception:
                pass
    elif wanted:
        note = "Maaf, gambar untuk itu belum tersedia ya."
        try:
            telegram_api.send_message(token, chat_id,
                                      f"{clean_reply}\n\n{note}" if clean_reply else note,
                                      reply_markup=pay_kb)
            image_note_sent = True
        except Exception:
            pass

    # customer minta qris & ada gambar QRIS -> kirim gambarnya
    if bot.qris_image_path and "qris" in text.lower():
        try:
            telegram_api.send_photo(token, chat_id, config.resolve_upload_path(bot.qris_image_path), caption="QRIS pembayaran")
            try:
                telegram_api.send_message(token, chat_id, "Itu gambar QRIS resmi dari penjual di atas ya. Setelah membayar, kirimkan foto/screenshot bukti pembayaran di chat ini.")
            except Exception:
                pass
        except Exception:
            telegram_api.send_message(token, chat_id, "Maaf, gambar QRIS gagal dikirim. Coba minta lagi atau hubungi penjual ya.")

    # ponytail: sertakan "incomplete" -> order yang masih dilengkapi tidak dobel
    existing = (db.query(ExtractedOrder)
                .join(Message, Message.id == ExtractedOrder.message_id)
                .filter(ExtractedOrder.bot_id == bot_id, Message.chat_id == str(chat_id),
                        ExtractedOrder.status.in_(["pending", "incomplete"]))
                .order_by(desc(ExtractedOrder.created_at)).first())
    existing_summary = None
    if existing:
        existing_summary = {
            "products": existing.products, "total_price": float(existing.total_price) if existing.total_price else None,
            "delivery_address": existing.delivery_address, "customer_phone": existing.customer_phone,
        }
    try:
        order = gemini_service.extract_order(api_key, text, history, existing_summary, (file_context or "") + "\n" + (bot.system_prompt or ""))
    except Exception:
        order = None

    msg = Message(bot_id=bot_id, user_id=uid, chat_id=str(chat_id), message_text=text,
                  response_text=reply, extracted_data=order, response_time=response_time,
                  model_used=model_used)
    db.add(msg)
    db.commit()
    db.refresh(msg)

    # ponytail: hanya simpan/update order setelah customer mengonfirmasi ringkasan dari AI
    saved_order = None
    if order and existing and order_is_confirmed(order):
        # 1 chat 1 open order — selalu merge ke order yang sama, field null tidak menimpa isi lama
        new_products = order.get("products") or existing.products
        # ponytail: note per produk — bila AI menghilangkan note lama untuk produk yg sama, pulihkan
        try:
            old_notes = {str(p.get("product_name", "")).lower(): p.get("note") for p in (existing.products or [])}
            for p in new_products:
                if not p.get("note"):
                    old = old_notes.get(str(p.get("product_name", "")).lower())
                    if old:
                        p["note"] = old
        except Exception:
            pass
        existing.products = new_products
        existing.total_price = order.get("total_price") if order.get("total_price") is not None else existing.total_price
        existing.delivery_address = order.get("delivery_address") or existing.delivery_address
        existing.customer_phone = order.get("customer_phone") or existing.customer_phone
        if not existing.customer_name and order.get("customer_name"):
            existing.customer_name = order.get("customer_name")
        # ponytail: status TIDAK diubah dari sisi chat — hanya creator yang boleh mengubah status order
        db.commit()
        saved_order = existing
    elif order and order_is_confirmed(order):
        status = "pending" if order.get("total_price") is not None else "incomplete"
        new_order = ExtractedOrder(
            bot_id=bot_id, message_id=msg.id, customer_user_id=uid,
            customer_name=order.get("customer_name") or customer_name,
            products=order["products"], total_price=order.get("total_price"),
            delivery_address=order.get("delivery_address"), customer_phone=order.get("customer_phone"),
        )
        db.add(new_order)
        db.commit()
        db.refresh(new_order)
        saved_order = new_order
        try:
            excel_service.append_order(bot_id, order, uid, status)
        except Exception:
            pass
        try:
            sheets = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot_id, BotExternalSource.kind == "sheet").all()
            for _p in (saved_order.products or []):
                for _s in sheets:
                    try:
                        if gsvc.adjust_stock(_s, str(_p.get("product_name", "")), -int(_p.get("quantity", 0) or 0)):
                            break
                    except Exception:
                        continue
            saved_order.stock_deducted = True
            db.commit()
        except Exception:
            pass

    if not image_note_sent:
        try:
            telegram_api.send_message(token, chat_id,
                                      clean_reply if clean_reply else "Berikut gambarnya ya.",
                                      reply_markup=pay_kb)
        except Exception:
            pass

    # data baru saja lengkap -> ingatkan bayar deterministik (balasan AI turn ini belum tahu)
    if saved_order and not data_complete and payment.payment_stage_active(bot, saved_order):
        try:
            text_r = payment.reminder_text(bot, saved_order)
            telegram_api.send_message(token, chat_id, text_r,
                                      reply_markup=payment.payment_keyboard(bot))
            db.add(Message(bot_id=bot_id, user_id=uid, chat_id=str(chat_id),
                           message_text="[pengingat pembayaran otomatis]", response_text=text_r))
            db.commit()
        except Exception:
            pass
    elif not saved_order and existing and not data_complete and order:
        # order belum terkonfirmasi, tapi ekstraksi turn ini menunjukkan data sudah lengkap
        merged = SimpleNamespace(
            customer_name=order.get("customer_name") or existing.customer_name,
            customer_phone=order.get("customer_phone") or existing.customer_phone,
            products=order.get("products") or existing.products,
            total_price=order.get("total_price") if order.get("total_price") is not None else existing.total_price,
            delivery_address=order.get("delivery_address") or existing.delivery_address,
            status=existing.status, payment_proof_path=existing.payment_proof_path)
        if payment.payment_stage_active(bot, merged):
            reminded = (db.query(Message)
                        .filter(Message.bot_id == bot_id, Message.chat_id == str(chat_id),
                                Message.response_text.like(payment.REMINDER_PREFIX + "%"))
                        .first())
            if not reminded:
                try:
                    text_r = payment.reminder_text(bot, merged)
                    telegram_api.send_message(token, chat_id, text_r,
                                              reply_markup=payment.payment_keyboard(bot))
                    db.add(Message(bot_id=bot_id, user_id=uid, chat_id=str(chat_id),
                                   message_text="[pengingat pembayaran otomatis]", response_text=text_r))
                    db.commit()
                except Exception:
                    pass

    emit_conversation(db, bot_id, uid, msg)
    return {"ok": True}


@app.on_event("startup")
def setup_webhook():
    if config.TELEGRAM_TOKEN and config.BASE_URL.startswith("https://"):
        try:
            telegram_api.set_webhook(config.TELEGRAM_TOKEN, f"{config.BASE_URL}/api/telegram/webhook", config.TELEGRAM_WEBHOOK_SECRET or None)
        except Exception:
            pass
    if config.TELEGRAM_TOKEN and not config.TELEGRAM_BOT_USERNAME:
        try:
            config.TELEGRAM_BOT_USERNAME = telegram_api.get_me(config.TELEGRAM_TOKEN).get("username", "")
        except Exception:
            pass


@app.get("/api/health")
def health():
    return {"status": "ok"}
