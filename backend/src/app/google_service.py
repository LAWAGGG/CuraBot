import os
import re
import threading
import time
from urllib.parse import urlparse
from . import config

ALLOWED_HOSTS = {"docs.google.com", "drive.google.com"}
ID_RE = re.compile(r"^[A-Za-z0-9-_]{10,}$")
_SHEET_PATTERNS = [re.compile(r"/spreadsheets/d/([A-Za-z0-9-_]+)"), re.compile(r"[?&]id=([A-Za-z0-9-_]{10,})")]
_FOLDER_PATTERNS = [re.compile(r"/folders/([A-Za-z0-9-_]+)"), re.compile(r"[?&]id=([A-Za-z0-9-_]{10,})")]
_SCOPES = ["https://www.googleapis.com/auth/spreadsheets", "https://www.googleapis.com/auth/drive.readonly"]

_locks: dict[int, threading.Lock] = {}
_locks_guard = threading.Lock()
_sheet_cache: dict[int, tuple[float, str]] = {}
_drive_cache: dict[int, tuple[float, list]] = {}


def _lock_for(bot_id: int) -> threading.Lock:
    with _locks_guard:
        return _locks.setdefault(bot_id, threading.Lock())


def parse_google_id(kind: str, url: str) -> str:
    u = (url or "").strip()
    if len(u) > 2000:
        raise ValueError("URL terlalu panjang")
    try:
        host = urlparse(u).hostname or ""
    except Exception:
        raise ValueError("URL tidak valid")
    if host not in ALLOWED_HOSTS:
        raise ValueError("Hanya link docs.google.com / drive.google.com")
    pats = _SHEET_PATTERNS if kind == "sheet" else _FOLDER_PATTERNS
    for p in pats:
        m = p.search(u)
        if m and ID_RE.match(m.group(1)):
            return m.group(1)
    raise ValueError("ID Google tidak ditemukan di link")


def _creds():
    from google.oauth2 import service_account
    path = config.GOOGLE_CREDENTIALS_PATH
    if not path or not os.path.isfile(path):
        raise RuntimeError("Kredensial Google belum dikonfigurasi (GOOGLE_CREDENTIALS_PATH)")
    return service_account.Credentials.from_service_account_file(path, scopes=_SCOPES)


def _drive_service():
    from googleapiclient.discovery import build
    return build("drive", "v3", credentials=_creds(), cache_discovery=False)


def guess_mapping(headers: list) -> dict:
    def find(*keys):
        for i, h in enumerate(headers):
            hl = re.sub(r"[^a-z0-9 ]", "", (h or "").lower())
            if any(k in hl for k in keys):
                return i
        return None
    return {
        "name_col": find("nama", "produk", "product", "item", "barang", "menu"),
        "price_col": find("harga", "price", "rp", "tarif"),
        "stock_col": find("stok", "stock", "sisa", "qty", "jumlah", "persediaan"),
        "image_col": find("gambar", "image", "foto", "drive"),
    }


def _norm(s: str) -> str:
    return re.sub(r"\s+", " ", (s or "").lower().strip())


def fuzzy_find(name: str, candidates: list) -> str | None:
    n = _norm(name)
    if not n:
        return None
    for c in candidates:
        if n and n in _norm(c):
            return c
    words = [w for w in re.findall(r"[a-z0-9]+", n) if len(w) > 3]
    if not words:
        return None
    for c in candidates:
        cl = _norm(c)
        if all(w in cl for w in words):
            return c
    return None


def read_sheet_rows(source) -> tuple:
    import gspread
    gc = gspread.service_account(filename=config.GOOGLE_CREDENTIALS_PATH)
    sh = gc.open_by_key(source.external_id)
    ws = sh.worksheet(source.tab) if source.tab else sh.sheet1
    values = ws.get_all_values()
    if not values:
        return [], [], guess_mapping([])
    headers = [(c or "").strip() for c in values[0]]
    rows = [r for r in values[1:] if any((c or "").strip() for c in r)]
    auto = guess_mapping(headers)
    saved = source.mapping or {}
    mapping = {**auto, **{k: v for k, v in saved.items() if v is not None}}
    return headers, rows, mapping


def sheet_context(bot_id: int, db) -> str:
    from .models import BotExternalSource
    now = time.time()
    hit = _sheet_cache.get(bot_id)
    if hit and now - hit[0] < config.GOOGLE_SHEET_TTL:
        return hit[1]
    lines: list[str] = []
    sources = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot_id, BotExternalSource.kind == "sheet").all()
    for s in sources:
        try:
            headers, rows, _m = read_sheet_rows(s)
            # ponytail: AI baca semua kolom apa adanya, tanpa mapping manual
            labels = [(h or f"Kolom{i + 1}").strip() for i, h in enumerate(headers)]
            for r in rows:
                parts = []
                for i, h in enumerate(labels):
                    v = (r[i].strip() if i < len(r) and r[i] else "")
                    if v:
                        parts.append(f"{h}: {v}"[:150])
                if not parts:
                    continue
                lines.append("- " + " | ".join(parts))
                if len("\n".join(lines)) > 7500:
                    break
            s.last_error = None
        except Exception as e:
            s.last_error = _safe_err(e)
    try:
        db.commit()
    except Exception:
        db.rollback()
    ctx = "Katalog Google Sheets:\n" + "\n".join(lines[:200]) if lines else ""
    _sheet_cache[bot_id] = (now, ctx)
    return ctx


IMAGE_PREFIX = "image/"
VIDEO_PREFIX = "video/"
GOOGLE_DOC_MIME = "application/vnd.google-apps.document"
GOOGLE_SLIDE_MIME = "application/vnd.google-apps.presentation"
GOOGLE_SHEET_MIME = "application/vnd.google-apps.spreadsheet"
_GOOGLE_EXPORT_MIMES = {GOOGLE_DOC_MIME: "text/plain", GOOGLE_SLIDE_MIME: "text/plain", GOOGLE_SHEET_MIME: "text/csv"}
_BINARY_DOC_MIMES = {
    "application/pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "application/msword",
    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "application/vnd.ms-excel",
    "text/plain",
    "text/csv",
}
_BINARY_DOC_EXTS = (".pdf", ".docx", ".doc", ".xlsx", ".xls", ".txt", ".csv")
_VIDEO_EXTS = (".mp4", ".mov", ".avi", ".mkv", ".webm", ".3gp", ".m4v")
MAX_DRIVE_DOC_BYTES = 10 * 1024 * 1024
MAX_DRIVE_VIDEO_BYTES = 50 * 1024 * 1024
_DRIVE_LIST_LIMIT = 500
_DRIVE_DOCS_PER_SOURCE = 20
_drive_docs_cache: dict[int, tuple[float, str]] = {}


def classify_mime(mime: str, name: str = "") -> str:
    """Klasifikasi murni tanpa I/O: image | doc | video | unsupported."""
    m = (mime or "").lower()
    if m.startswith(IMAGE_PREFIX):
        return "image"
    if m.startswith(VIDEO_PREFIX):
        return "video"
    if m in _GOOGLE_EXPORT_MIMES or m in _BINARY_DOC_MIMES:
        return "doc"
    ext = os.path.splitext(name or "")[1].lower()
    if ext in _BINARY_DOC_EXTS:
        return "doc"
    if ext in _VIDEO_EXTS:
        return "video"
    return "unsupported"


def _list_drive_all(source) -> list:
    now = time.time()
    hit = _drive_cache.get(source.id)
    if hit and now - hit[0] < config.GOOGLE_DRIVE_TTL:
        return hit[1]
    svc = _drive_service()
    # ponytail: flat, tanpa rekursi subfolder/shortcut -> anti traversal keluar folder
    out, token = [], None
    while True:
        resp = svc.files().list(q=f"'{source.external_id}' in parents and trashed=false", fields="files(id,name,mimeType),nextPageToken", pageSize=100, pageToken=token).execute(num_retries=2)
        for f in resp.get("files", []):
            if f.get("id") and f.get("name"):
                out.append({"id": f["id"], "name": f.get("name", ""), "mime": f.get("mimeType", "")})
            if len(out) >= _DRIVE_LIST_LIMIT:
                break
        token = resp.get("nextPageToken")
        if not token or len(out) >= _DRIVE_LIST_LIMIT:
            break
    _drive_cache[source.id] = (now, out)
    return out


def list_drive_images(source) -> list:
    return [{"id": f["id"], "name": f["name"]} for f in _list_drive_all(source) if classify_mime(f.get("mime", ""), f.get("name", "")) == "image"]


def list_drive_docs(source) -> list:
    return [f for f in _list_drive_all(source) if classify_mime(f.get("mime", ""), f.get("name", "")) == "doc"]


def list_drive_videos(source) -> list:
    return [f for f in _list_drive_all(source) if classify_mime(f.get("mime", ""), f.get("name", "")) == "video"]


def drive_summary(source) -> dict:
    """Ringkasan folder campur untuk sync/preview. Tanpa download isi file."""
    images, docs, videos = [], [], []
    unsupported = 0
    for f in _list_drive_all(source):
        kind = classify_mime(f.get("mime", ""), f.get("name", ""))
        if kind == "image":
            images.append(f["name"])
        elif kind == "doc":
            docs.append(f["name"])
        elif kind == "video":
            videos.append(f["name"])
        else:
            unsupported += 1
    return {
        "images": len(images), "docs": len(docs), "videos": len(videos),
        "unsupported": unsupported, "preview": images[:5],
        "preview_docs": docs[:5], "preview_videos": videos[:5],
    }


def _download_request(req, max_bytes: int) -> bytes:
    from googleapiclient.http import MediaIoBaseDownload
    import io as _io
    buf = _io.BytesIO()
    dl = MediaIoBaseDownload(buf, req)
    done = False
    while not done:
        _, done = dl.next_chunk()
    data = buf.getvalue()
    if not data or len(data) > max_bytes:
        raise ValueError("Ukuran file Drive tidak valid")
    return data


def download_drive_file(file_id: str, max_bytes: int = 25 * 1024 * 1024) -> bytes:
    if not ID_RE.match(file_id or ""):
        raise ValueError("file_id tidak valid")
    svc = _drive_service()
    return _download_request(svc.files().get_media(fileId=file_id), max_bytes)


def download_drive_image(file_id: str) -> bytes:
    if not ID_RE.match(file_id or ""):
        raise ValueError("file_id tidak valid")
    try:
        return download_drive_file(file_id)
    except ValueError:
        raise ValueError("Ukuran gambar Drive tidak valid")


def export_drive_text(file_id: str, mime: str) -> bytes:
    if not ID_RE.match(file_id or ""):
        raise ValueError("file_id tidak valid")
    export_mime = _GOOGLE_EXPORT_MIMES.get((mime or "").lower())
    if not export_mime:
        raise ValueError("bukan dokumen Google")
    svc = _drive_service()
    return _download_request(svc.files().export_media(fileId=file_id, mimeType=export_mime), MAX_DRIVE_DOC_BYTES)


def extract_drive_text(name: str, mime: str, data: bytes) -> str:
    """Parse bytes -> teks. Sama seperti parser upload (pdf/docx/xlsx/txt/csv)."""
    if not data:
        return ""
    m = (mime or "").lower()
    if m in _GOOGLE_EXPORT_MIMES or m in ("text/plain", "text/csv"):
        try:
            return data.decode("utf-8", errors="ignore").strip()[:8000]
        except Exception:
            return ""
    ext = os.path.splitext(name or "")[1].lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            import io as _io
            reader = PdfReader(_io.BytesIO(data))
            return "\n".join((p.extract_text() or "") for p in reader.pages).strip()[:8000]
        if ext in (".docx", ".doc"):
            from docx import Document
            import io as _io
            doc = Document(_io.BytesIO(data))
            return "\n".join(p.text for p in doc.paragraphs).strip()[:8000]
        if ext in (".xlsx", ".xls"):
            from openpyxl import load_workbook
            import io as _io
            wb = load_workbook(_io.BytesIO(data), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets:
                out.append(f"[Sheet: {ws.title}]")
                for row in ws.iter_rows(values_only=True):
                    cells = ["" if c is None else str(c) for c in row]
                    if any(cells):
                        out.append(" | ".join(cells))
            wb.close()
            return "\n".join(out).strip()[:8000]
        if ext in (".txt", ".csv"):
            return data.decode("utf-8", errors="ignore").strip()[:8000]
    except Exception:
        return ""
    return ""


def drive_docs_context(bot_id: int, db) -> str:
    from .models import BotExternalSource
    now = time.time()
    hit = _drive_docs_cache.get(bot_id)
    if hit and now - hit[0] < config.GOOGLE_DRIVE_TTL:
        return hit[1]
    chunks: list[str] = []
    try:
        sources = db.query(BotExternalSource).filter(BotExternalSource.bot_id == bot_id, BotExternalSource.kind == "drive_folder").all()
    except Exception:
        sources = []
    for s in sources:
        try:
            docs = list_drive_docs(s)
        except Exception:
            continue
        for d in docs[:_DRIVE_DOCS_PER_SOURCE]:
            try:
                mime = d.get("mime", "")
                if mime.lower() in _GOOGLE_EXPORT_MIMES:
                    raw = export_drive_text(d["id"], mime)
                else:
                    raw = download_drive_file(d["id"], MAX_DRIVE_DOC_BYTES)
                text = extract_drive_text(d.get("name", ""), mime, raw).strip()
            except Exception:
                continue
            if not text:
                continue
            chunks.append(f"[{d.get('name', '')}]\n{text[:2000]}")
            if len("\n\n".join(chunks)) > 4000:
                break
        if len("\n\n".join(chunks)) > 4000:
            break
    ctx = "Dokumen Google Drive:\n" + "\n\n".join(chunks) if chunks else ""
    _drive_docs_cache[bot_id] = (now, ctx)
    return ctx


def adjust_stock(source, product_name: str, delta: int) -> bool:
    import gspread
    with _lock_for(source.bot_id):
        gc = gspread.service_account(filename=config.GOOGLE_CREDENTIALS_PATH)
        sh = gc.open_by_key(source.external_id)
        ws = sh.worksheet(source.tab) if source.tab else sh.sheet1
        values = ws.get_all_values()
        if len(values) < 2:
            return False
        headers = values[0]
        m = {**guess_mapping(headers), **{k: v for k, v in (source.mapping or {}).items() if v is not None}}
        nc, sc = m.get("name_col"), m.get("stock_col")
        if nc is None or sc is None:
            return False
        names = [((r[nc] if nc < len(r) else "") or "") for r in values[1:]]
        target = fuzzy_find(product_name, names)
        if not target:
            return False
        idx = names.index(target) + 2
        try:
            cur = int(str(values[idx - 1][sc]).strip() or "0")
        except (ValueError, IndexError):
            return False
        ws.update_cell(idx, sc + 1, str(max(0, cur + delta)))
        _sheet_cache.pop(source.bot_id, None)
        return True


def _safe_err(e: Exception) -> str:
    msg = f"{type(e).__name__}: {e}"
    for secret in (config.GOOGLE_CREDENTIALS_PATH,):
        if secret:
            msg = msg.replace(secret, "[redacted]")
    return msg[:500]
